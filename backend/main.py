from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from pathlib import Path
from copy import copy
import pandas as pd
import json
import re
import openpyxl


app = FastAPI()


app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


BASE_DIR = Path(__file__).resolve().parent


PASTA_PLANILHAS = BASE_DIR.parent / "planilhas"

NOME_LEVANTAMENTO = "levantamento.xlsx"

ARQUIVO_ESTADO = BASE_DIR / "estado.json"


class Conclusao(BaseModel):
    peg: str
    quantidade_caixa: int


class Andamento(BaseModel):
    pegs: list[str]


class CancelarAndamento(BaseModel):
    pegs: list[str]

def carregar_estado():
    if not ARQUIVO_ESTADO.exists():
        return []

    try:
        with open(
            ARQUIVO_ESTADO,
            "r",
            encoding="utf-8"
        ) as arquivo:
            dados = json.load(arquivo)

        if not isinstance(dados, list):
            return []

        return consolidar_pegs(dados)

    except Exception:
        return []

def salvar_estado(dados):
    dados = consolidar_pegs(dados)

    with open(
        ARQUIVO_ESTADO,
        "w",
        encoding="utf-8"
    ) as arquivo:
        json.dump(
            dados,
            arquivo,
            ensure_ascii=False,
            indent=2
        )

def consolidar_pegs(dados):

    mapa = {}

    prioridade = {
        "pendente": 1,
        "andamento": 2,
        "concluido": 3
    }

    for item in dados:
        peg = normalizar_peg(item.get("peg"))

        if not peg:
            continue

        item["peg"] = peg

        existente = mapa.get(peg)

        if existente is None:
            mapa[peg] = item
            continue

        status_atual = existente.get("status", "pendente")
        status_novo = item.get("status", "pendente")

        if prioridade.get(status_novo, 0) > prioridade.get(status_atual, 0):
            mapa[peg] = item

        elif (
            status_novo == "concluido"
            and item.get("quantidade_caixa") is not None
        ):
            mapa[peg] = item

    return list(mapa.values())


def normalizar_peg(valor):
    if pd.isna(valor):
        return ""

    texto = str(valor).strip().upper()

    if re.fullmatch(r"\d+\.0", texto):
        texto = texto[:-2]

    texto = texto.replace(".", "")

    return texto



def encontrar_planilha_estoque():

    arquivos = list(
        PASTA_PLANILHAS.glob("*.xlsx")
    )

    arquivos = [
        arquivo
        for arquivo in arquivos
        if arquivo.name.lower()
        != NOME_LEVANTAMENTO.lower()
    ]

    if not arquivos:
        raise ValueError(
            "Nenhuma planilha de estoque encontrada."
        )

    def extrair_data(arquivo):

        nome = arquivo.stem

        numeros = re.findall(
            r"\d+",
            nome
        )

        if len(numeros) >= 3:

            try:

                dia = int(numeros[0])
                mes = int(numeros[1])
                ano = int(numeros[2])

                if ano < 100:
                    ano += 2000

                return pd.Timestamp(
                    year=ano,
                    month=mes,
                    day=dia
                )

            except Exception:
                pass

        return pd.Timestamp(
            arquivo.stat().st_mtime,
            unit="s"
        )

    arquivos.sort(
        key=extrair_data,
        reverse=True
    )

    return arquivos[0]


def processar_estoque(path):

    df = pd.read_excel(
        path,
        engine="openpyxl"
    )

    colunas = list(
        df.columns
    )

    if len(colunas) < 11:
        raise ValueError(
            "A planilha de estoque precisa ter pelo menos 11 colunas."
        )

    localizacao_col = colunas[3]
    peg_col = colunas[7]
    descricao_col = colunas[10]

    dados = df[
        [
            localizacao_col,
            peg_col,
            descricao_col
        ]
    ].copy()

    dados.columns = [
        "localizacao",
        "peg",
        "descricao"
    ]

    dados = dados.dropna(
        subset=["peg"]
    )

    dados["localizacao"] = (
        dados["localizacao"]
        .astype(str)
        .str.strip()
    )

    dados["peg"] = (
        dados["peg"]
        .apply(normalizar_peg)
    )

    dados["descricao"] = (
        dados["descricao"]
        .astype(str)
        .str.strip()
    )

    dados = dados[
        ~dados["descricao"]
        .str.upper()
        .str.startswith(
            "PAPEL",
            na=False
        )
    ].copy()

    dados = dados[
        dados["localizacao"].str.startswith(
            ("CDA", "CDL"),
            na=False
        )
    ].copy()

    dados["numero_local"] = pd.to_numeric(
        dados["localizacao"]
        .str.extract(
            r"(?:CDA|CDL)\s*(\d+)",
            expand=False
        ),
        errors="coerce"
    )

    dados = dados[
        ~dados["localizacao"].str.startswith(
            "CDA"
        )
        |
        (
            dados["localizacao"].str.contains(
                "Y01",
                na=False
            )
            &
            (
                dados["numero_local"] <= 15
            )
        )
    ].copy()

    dados = dados[
        ~dados["localizacao"].str.startswith(
            "CDL"
        )
        |
        (
            dados["numero_local"] <= 44
        )
    ].copy()

    resultado = (
        dados
        .groupby(
            "peg",
            as_index=False
        )
        .agg(
            descricao=(
                "descricao",
                "first"
            ),
            quantidade_itens=(
                "peg",
                "size"
            ),
            localizacoes=(
                "localizacao",
                lambda x: ", ".join(
                    sorted(set(x))
                )
            )
        )
    )

    return resultado

def processar_levantamento(path):
    df = pd.read_excel(
        path,
        engine="openpyxl",
        dtype=object
    )

    if len(df.columns) < 3:
        raise ValueError(
            "A planilha de levantamento precisa ter pelo menos 3 colunas."
        )

    df = df.iloc[:, :3].copy()

    df.columns = [
        "peg",
        "descricao",
        "quantidade_caixa"
    ]

    df = df.dropna(
        subset=["peg"]
    )

    df["peg"] = df["peg"].apply(
        normalizar_peg
    )

    df["descricao"] = (
        df["descricao"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    df["quantidade_caixa"] = pd.to_numeric(
        df["quantidade_caixa"],
        errors="coerce"
    )

    df = df[
        df["quantidade_caixa"].notna()
        &
        (df["quantidade_caixa"] > 0)
    ].copy()

    df = df.drop_duplicates(
        subset=["peg"],
        keep="last"
    )

    return df

def aplicar_status_antigo(dados_novos):

    estado_antigo = carregar_estado()

    mapa_antigo = {
        item["peg"]: item
        for item in estado_antigo
    }

    for item in dados_novos:

        antigo = mapa_antigo.get(
            item["peg"]
        )

        if antigo:

            if (
                antigo.get(
                    "quantidade_caixa"
                ) is not None
                and
                antigo.get(
                    "quantidade_caixa"
                ) > 0
            ):

                item["quantidade_caixa"] = (
                    antigo[
                        "quantidade_caixa"
                    ]
                )

                item["status"] = "concluido"

                continue

            if antigo.get(
                "status"
            ) == "andamento":

                item["status"] = "andamento"

                continue

        item["status"] = "pendente"

    return dados_novos


@app.post("/atualizar")
def atualizar_planilhas():

    try:

        if not PASTA_PLANILHAS.exists():
            raise ValueError(
                f"A pasta planilhas não existe: {PASTA_PLANILHAS}"
            )

        estoque_path = (
            encontrar_planilha_estoque()
        )

        levantamento_path = (
            PASTA_PLANILHAS /
            NOME_LEVANTAMENTO
        )

        if not levantamento_path.exists():
            raise ValueError(
                f"Planilha de levantamento não encontrada: "
                f"{levantamento_path}"
            )

        print(
            "\n=============================="
        )

        print(
            "ATUALIZANDO PLANILHAS"
        )

        print(
            "=============================="
        )

        print(
            f"Estoque: {estoque_path.name}"
        )

        print(
            f"Levantamento: "
            f"{levantamento_path.name}"
        )

        estoque = processar_estoque(
            estoque_path
        )

        levantamento = processar_levantamento(
            levantamento_path
        )

        print(
            f"PEGs encontrados no estoque: "
            f"{len(estoque)}"
        )

        print(
            f"PEGs concluídos no levantamento: "
            f"{len(levantamento)}"
        )

        print(
            "\n--- EXEMPLO DE PEGs DO ESTOQUE ---"
        )

        print(
            estoque["peg"].head(10).tolist()
        )

        print(
            "\n--- EXEMPLO DE CÓDIGOS DO LEVANTAMENTO ---"
        )

        print(
            levantamento["peg"].head(10).tolist()
        )
        estoque["peg"] = (
            estoque["peg"]
            .apply(normalizar_peg)
        )

        levantamento["peg"] = (
            levantamento["peg"]
            .apply(normalizar_peg)
        )

        quantidades = dict(
            zip(
                levantamento["peg"],
                levantamento[
                    "quantidade_caixa"
                ]
            )
        )

        dados = []

        concluidos_importados = 0

        for item in estoque.to_dict(
            orient="records"
        ):

            peg = normalizar_peg(
                item["peg"]
            )

            quantidade = quantidades.get(
                peg
            )

            if quantidade is not None:
                concluidos_importados += 1

            dados.append({

                "peg": peg,

                "descricao": item[
                    "descricao"
                ],

                "quantidade_itens": item[
                    "quantidade_itens"
                ],

                "localizacoes": item[
                    "localizacoes"
                ],

                "quantidade_caixa": (
                    int(quantidade)
                    if pd.notna(
                        quantidade
                    )
                    else None
                ),

                "status": (
                    "concluido"
                    if quantidade is not None
                    else "pendente"
                )

            })

        dados = aplicar_status_antigo(
            dados
        )

        for item in dados:

            quantidade = quantidades.get(
                item["peg"]
            )

            if quantidade is not None:

                item["quantidade_caixa"] = (
                    int(quantidade)
                )

                item["status"] = "concluido"

        print(
            f"PEGs concluídos cruzados: "
            f"{concluidos_importados}"
        )

        salvar_estado(
            dados
        )

        return {

            "mensagem":
                "Planilhas atualizadas com sucesso.",

            "estoque":
                estoque_path.name,

            "levantamento":
                levantamento_path.name,

            "total_pegs":
                len(dados),

            "concluidos_importados":
                concluidos_importados,

            "dados":
                dados

        }

    except Exception as erro:

        print(
            f"\nERRO AO ATUALIZAR: {erro}"
        )

        return {
            "erro": str(erro)
        }


@app.get("/dados")
def obter_dados():

    return {
        "dados":
            carregar_estado()
    }


@app.post("/iniciar-andamento")
def iniciar_andamento(andamento: Andamento):
    try:
        estado = carregar_estado()

        pegs_selecionados = {
            normalizar_peg(peg)
            for peg in andamento.pegs
        }

        enviados = 0
        ignorados = []

        for item in estado:
            peg_atual = normalizar_peg(item["peg"])

            if peg_atual not in pegs_selecionados:
                continue

            if item.get("status") == "pendente":
                item["status"] = "andamento"
                enviados += 1
            else:
                ignorados.append({
                    "peg": item["peg"],
                    "status": item.get("status")
                })

        salvar_estado(estado)

        return {
            "mensagem": f"{enviados} PEG(s) enviados para andamento.",
            "enviados": enviados,
            "ignorados": ignorados,
            "dados": estado
        }

    except Exception as erro:
        return {"erro": str(erro)}

def normalizar_codigo_planilha(valor):

    if valor is None:
        return ""

    texto = str(valor).strip().upper()

    texto = re.sub(r"[^0-9A-Z]", "", texto)

    return texto

def atualizar_levantamento(peg, descricao, quantidade_caixa):
    caminho = (
        PASTA_PLANILHAS
        / NOME_LEVANTAMENTO
    ).resolve()

    if not caminho.exists():
        raise FileNotFoundError(
            f"Planilha de levantamento não encontrada: {caminho}"
        )

    caminho_backup = caminho.with_name(
        "Levantamento Padrão de Caixas Fechadas - Backup automático.xlsx"
    )

    if not caminho_backup.exists():
        import shutil

        shutil.copy2(
            caminho,
            caminho_backup
        )

    workbook = openpyxl.load_workbook(
        caminho
    )

    worksheet = workbook[
        workbook.sheetnames[0]
    ]

    codigo_alvo = normalizar_peg(
        peg
    )

    coluna_codigo = None
    coluna_descricao = None
    coluna_padrao = None

    for coluna in range(
        1,
        worksheet.max_column + 1
    ):
        valor = worksheet.cell(
            row=1,
            column=coluna
        ).value

        if valor is None:
            continue

        cabecalho = (
            str(valor)
            .strip()
            .upper()
        )

        if cabecalho in (
            "CÓDIGO",
            "CODIGO"
        ):
            coluna_codigo = coluna

        elif cabecalho == "DESCRIÇÃO":
            coluna_descricao = coluna

        elif cabecalho in (
            "PADRÃO EMBALAGEM",
            "PADRAO EMBALAGEM"
        ):
            coluna_padrao = coluna

    if coluna_codigo is None:
        workbook.close()

        raise ValueError(
            "A coluna 'Código' não foi encontrada."
        )

    if coluna_descricao is None:
        workbook.close()

        raise ValueError(
            "A coluna 'Descrição' não foi encontrada."
        )

    if coluna_padrao is None:
        workbook.close()

        raise ValueError(
            "A coluna 'Padrão Embalagem' não foi encontrada."
        )

    linha_encontrada = None

    for linha in range(
        2,
        worksheet.max_row + 1
    ):
        valor_codigo = worksheet.cell(
            row=linha,
            column=coluna_codigo
        ).value

        codigo_atual = normalizar_peg(
            valor_codigo
        )

        if codigo_atual == codigo_alvo:
            linha_encontrada = linha
            break

    if linha_encontrada is not None:

        worksheet.cell(
            row=linha_encontrada,
            column=coluna_descricao
        ).value = descricao

        worksheet.cell(
            row=linha_encontrada,
            column=coluna_padrao
        ).value = quantidade_caixa

        mensagem = (
            f"PEG {peg} já existia e foi atualizado."
        )

    else:

        nova_linha = worksheet.max_row + 1
        linha_modelo = worksheet.max_row

        if linha_modelo in worksheet.row_dimensions:
            worksheet.row_dimensions[nova_linha].height = (
                worksheet.row_dimensions[linha_modelo].height
            )

        for coluna in range(
            1,
            worksheet.max_column + 1
        ):
            celula_origem = worksheet.cell(
                row=linha_modelo,
                column=coluna
            )

            celula_destino = worksheet.cell(
                row=nova_linha,
                column=coluna
            )

            if celula_origem.has_style:
                celula_destino._style = copy(
                    celula_origem._style
                )

            if celula_origem.number_format:
                celula_destino.number_format = (
                    celula_origem.number_format
                )

            if celula_origem.alignment:
                celula_destino.alignment = copy(
                    celula_origem.alignment
                )

            if celula_origem.protection:
                celula_destino.protection = copy(
                    celula_origem.protection
                )

        worksheet.cell(
            row=nova_linha,
            column=coluna_codigo
        ).value = codigo_alvo

        worksheet.cell(
            row=nova_linha,
            column=coluna_descricao
        ).value = descricao

        worksheet.cell(
            row=nova_linha,
            column=coluna_padrao
        ).value = quantidade_caixa

        linha_encontrada = nova_linha

        mensagem = (
            f"PEG {peg} foi adicionado ao levantamento."
        )

    arquivo_temporario = (
        PASTA_PLANILHAS
        / "~levantamento_atualizando.xlsx"
    )

    workbook.save(
        arquivo_temporario
    )

    workbook.close()

    teste = openpyxl.load_workbook(
        arquivo_temporario,
        read_only=True
    )

    teste.close()

    arquivo_temporario.replace(
        caminho
    )

    return {
        "peg": peg,
        "quantidade_caixa": quantidade_caixa,
        "linha": linha_encontrada,
        "novo": mensagem.endswith(
            "foi adicionado ao levantamento."
        ),
        "mensagem": mensagem
    }

@app.post("/concluir-peg")
def concluir_peg(conclusao: Conclusao):
    try:
        estado = carregar_estado()

        peg_alvo = normalizar_peg(conclusao.peg)

        if conclusao.quantidade_caixa <= 0:
            return {
                "erro": "A quantidade por caixa precisa ser maior que zero."
            }

        encontrado = False
        descricao = ""

        for item in estado:
            peg_atual = normalizar_peg(item["peg"])

            if peg_atual != peg_alvo:
                continue

            encontrado = True

            if item.get("status") != "andamento":
                return {
                    "erro": (
                        f"O PEG {conclusao.peg} não pode ser concluído "
                        f"porque está com status '{item.get('status')}'."
                    )
                }

            descricao = item.get("descricao", "")

            item["quantidade_caixa"] = conclusao.quantidade_caixa
            item["status"] = "concluido"

            break

        if not encontrado:
            return {
                "erro": f"PEG {conclusao.peg} não encontrado."
            }

        resultado_planilha = atualizar_levantamento(
            peg_alvo,
            descricao,
            conclusao.quantidade_caixa
        )

        salvar_estado(estado)

        return {
            "mensagem": (
                f"PEG {conclusao.peg} concluído "
                "e planilha atualizada."
            ),
            "dados": estado,
            "planilha": resultado_planilha
        }

    except Exception as erro:
        return {
            "erro": str(erro)
        }

@app.get("/")
def inicio():

    return {
        "mensagem":
            "API funcionando"
    }


@app.post("/cancelar-andamento")
def cancelar_andamento(
    cancelamento: CancelarAndamento
):

    try:

        estado = carregar_estado()

        pegs_selecionados = set(
            cancelamento.pegs
        )

        encontrados = 0

        for item in estado:

            if (
                item["peg"]
                in pegs_selecionados
                and
                item.get("status")
                == "andamento"
            ):

                item["status"] = "pendente"

                item["quantidade_caixa"] = None

                encontrados += 1

        salvar_estado(
            estado
        )

        return {
            "mensagem":
                f"{encontrados} PEG(s) voltaram para pendências.",

            "dados":
                estado
        }

    except Exception as erro:

        return {
            "erro": str(erro)
        }