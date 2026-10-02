from copy import copy
from pathlib import Path
import re
import shutil

import openpyxl
import pandas as pd

from services.utils import normalizar_peg


def encontrar_planilha_estoque(
    pasta_planilhas: Path,
    nome_levantamento: str
):
    arquivos = list(pasta_planilhas.glob("*.xlsx"))

    arquivos = [
        arquivo
        for arquivo in arquivos
        if (
            not arquivo.name.startswith("~")
            and arquivo.name.lower() != nome_levantamento.lower()
        )
    ]

    if not arquivos:
        raise ValueError(
            "Nenhuma planilha de estoque encontrada."
        )

    def extrair_data(arquivo):
        nome = arquivo.stem
        numeros = re.findall(r"\d+", nome)

        if len(numeros) >= 3:
            try:
                dia = int(numeros[0])
                mes = int(numeros[1])
                ano = int(numeros[2])

                if ano < 100:
                    ano += 2000

                return pd.Timestamp(
                    year=ano,
                    month=mes,
                    day=dia
                )

            except Exception:
                pass

        return pd.Timestamp(
            arquivo.stat().st_mtime,
            unit="s"
        )

    arquivos.sort(
        key=extrair_data,
        reverse=True
    )

    return arquivos[0]


def processar_estoque(path: Path):
    df = pd.read_excel(
        path,
        engine="openpyxl"
    )

    if df.shape[1] < 11:
        raise ValueError(
            f"A planilha de estoque foi lida com apenas "
            f"{df.shape[1]} colunas."
        )

    # Estrutura fixa da planilha de estoque:
    # D = Localização
    # H = PEG
    # K = Descrição
    localizacao_col = df.columns[3]
    peg_col = df.columns[7]
    descricao_col = df.columns[10]

    dados = df[
        [
            localizacao_col,
            peg_col,
            descricao_col
        ]
    ].copy()

    dados.columns = [
        "localizacao",
        "peg",
        "descricao"
    ]

    dados = dados.dropna(
        subset=["peg"]
    )

    dados["localizacao"] = (
        dados["localizacao"]
        .astype(str)
        .str.strip()
    )

    dados["peg"] = (
        dados["peg"]
        .apply(normalizar_peg)
    )

    dados["descricao"] = (
        dados["descricao"]
        .astype(str)
        .str.strip()
    )

    dados = dados[
        ~dados["descricao"]
        .str.upper()
        .str.startswith("PAPEL", na=False)
    ].copy()

    dados = dados[
        dados["localizacao"]
        .str.startswith(
            ("CDA", "CDL"),
            na=False
        )
    ].copy()

    dados["numero_local"] = pd.to_numeric(
        dados["localizacao"]
        .str.extract(
            r"(?:CDA|CDL)\s*(\d+)",
            expand=False
        ),
        errors="coerce"
    )

    dados = dados[
        ~dados["localizacao"].str.startswith("CDA")
        |
        (
            dados["localizacao"].str.contains(
                "Y01",
                na=False
            )
            & (dados["numero_local"] <= 15)
        )
    ].copy()

    dados = dados[
        ~dados["localizacao"].str.startswith("CDL")
        |
        (dados["numero_local"] <= 44)
    ].copy()

    resultado = (
        dados
        .groupby("peg", as_index=False)
        .agg(
            descricao=("descricao", "first"),
            quantidade_itens=("peg", "size"),
            localizacoes=(
                "localizacao",
                lambda x: ", ".join(
                    sorted(set(x))
                )
            )
        )
    )

    return resultado

def preencher_descricoes_levantamento(
    caminho_levantamento: Path,
    caminho_estoque: Path
):
    df_estoque = pd.read_excel(
        caminho_estoque,
        engine="openpyxl",
        dtype=object
    )

    if df_estoque.shape[1] < 11:
        raise ValueError(
            f"A planilha de estoque foi lida com apenas "
            f"{df_estoque.shape[1]} colunas."
        )

    peg_col = df_estoque.columns[7]
    descricao_col = df_estoque.columns[10]

    mapa_descricoes = {}

    for _, linha in df_estoque.iterrows():
        peg = normalizar_peg(
            linha[peg_col]
        )

        if not peg:
            continue

        descricao = linha[descricao_col]

        if pd.isna(descricao):
            continue

        descricao = str(
            descricao
        ).strip()

        if not descricao:
            continue

        if peg not in mapa_descricoes:
            mapa_descricoes[peg] = descricao

    workbook = openpyxl.load_workbook(
        caminho_levantamento
    )

    caminho_temporario = (
        caminho_levantamento
        .with_name("~levantamento_descricoes.xlsx")
    )

    alteracoes = 0

    try:
        sheet = workbook.active

        for linha in range(
            2,
            sheet.max_row + 1
        ):
            valor_peg = sheet.cell(
                linha,
                1
            ).value

            if valor_peg is None:
                continue

            peg = normalizar_peg(
                valor_peg
            )

            if not peg:
                continue

            descricao_atual = sheet.cell(
                linha,
                2
            ).value

            descricao_vazia = (
                descricao_atual is None
                or str(descricao_atual).strip() == ""
            )

            if not descricao_vazia:
                continue

            descricao_estoque = mapa_descricoes.get(
                peg
            )

            if descricao_estoque:
                sheet.cell(
                    linha,
                    2
                ).value = descricao_estoque

                alteracoes += 1

        workbook.save(
            caminho_temporario
        )

    finally:
        workbook.close()

    workbook_validacao = openpyxl.load_workbook(
        caminho_temporario,
        read_only=True
    )
    workbook_validacao.close()

    shutil.copy2(
        caminho_temporario,
        caminho_levantamento
    )

    caminho_temporario.unlink(
        missing_ok=True
    )

    return alteracoes

def processar_levantamento(
    path: Path,
    estoque: pd.DataFrame
):
    df = pd.read_excel(
        path,
        engine="openpyxl",
        dtype=object
    )

    if len(df.columns) < 3:
        raise ValueError(
            "A planilha de levantamento precisa ter pelo menos 3 colunas."
        )

    df = df.iloc[:, :3].copy()

    df.columns = [
        "peg",
        "descricao",
        "quantidade_caixa"
    ]

    df = df.dropna(
        subset=["peg"]
    )

    df["peg"] = (
        df["peg"]
        .apply(normalizar_peg)
    )

    df["descricao"] = (
        df["descricao"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    mapa_descricoes = (
        estoque[
            estoque["descricao"].notna()
            & (
                estoque["descricao"]
                .astype(str)
                .str.strip()
                != ""
            )
        ]
        .drop_duplicates(
            subset=["peg"]
        )
        .set_index("peg")["descricao"]
        .to_dict()
    )

    df["descricao"] = df.apply(
        lambda linha: (
            mapa_descricoes.get(
                linha["peg"],
                linha["descricao"]
            )
            if not linha["descricao"]
            else linha["descricao"]
        ),
        axis=1
    )

    df["quantidade_caixa"] = pd.to_numeric(
        df["quantidade_caixa"],
        errors="coerce"
    )

    df = df[
        df["quantidade_caixa"].notna()
        &
        (df["quantidade_caixa"] > 0)
    ].copy()

    df = df.drop_duplicates(
        subset=["peg"],
        keep="last"
    )

    return df


def atualizar_levantamento(
    caminho: Path,
    peg: str,
    descricao: str,
    quantidade_caixa: int,
    nome_backup: str,
):
    caminho_backup = caminho.with_name(
        nome_backup
    )

    if not caminho_backup.exists():
        shutil.copy2(
            caminho,
            caminho_backup
        )

    workbook = openpyxl.load_workbook(
        caminho
    )

    try:
        sheet = workbook.active

        # Estrutura fixa da planilha:
        # A = PEG
        # B = Descrição
        # C = Quantidade por caixa
        coluna_peg = 1
        coluna_descricao = 2
        coluna_quantidade = 3

        linha_encontrada = None

        # Primeiro procura se o PEG já existe.
        for linha in range(
            2,
            sheet.max_row + 1
        ):
            valor = sheet.cell(
                linha,
                coluna_peg
            ).value

            if (
                normalizar_peg(valor)
                == normalizar_peg(peg)
            ):
                linha_encontrada = linha
                break

        nova_linha = False

        if linha_encontrada is None:
            # Procura a primeira linha realmente vazia
            # na coluna A, ignorando a formatação das linhas abaixo.
            for linha in range(
                2,
                sheet.max_row + 1
            ):
                valor = sheet.cell(
                    linha,
                    coluna_peg
                ).value

                if valor is None or str(valor).strip() == "":
                    linha_encontrada = linha
                    break

            # Se não houver nenhuma linha vazia dentro
            # da área existente, cria uma nova linha.
            if linha_encontrada is None:
                linha_encontrada = sheet.max_row + 1

            nova_linha = True

            # Copia a formatação da linha anterior
            # para manter o padrão visual da planilha.
            if linha_encontrada > 2:
                linha_modelo = linha_encontrada - 1

                for coluna in range(
                    1,
                    sheet.max_column + 1
                ):
                    origem = sheet.cell(
                        linha_modelo,
                        coluna
                    )

                    destino = sheet.cell(
                        linha_encontrada,
                        coluna
                    )

                    if origem.has_style:
                        destino._style = copy(
                            origem._style
                        )

                    if origem.number_format:
                        destino.number_format = (
                            origem.number_format
                        )

                    if origem.alignment:
                        destino.alignment = copy(
                            origem.alignment
                        )

                    if origem.protection:
                        destino.protection = copy(
                            origem.protection
                        )

        sheet.cell(
            linha_encontrada,
            coluna_peg
        ).value = peg

        sheet.cell(
            linha_encontrada,
            coluna_descricao
        ).value = descricao

        sheet.cell(
            linha_encontrada,
            coluna_quantidade
        ).value = quantidade_caixa

        caminho_temporario = caminho.with_name(
            "~levantamento_atualizando.xlsx"
        )

        workbook.save(
            caminho_temporario
        )

    finally:
        workbook.close()

    # Valida o arquivo temporário antes de substituir
    # o arquivo original.
    workbook_validacao = openpyxl.load_workbook(
        caminho_temporario,
        read_only=True
    )

    workbook_validacao.close()

    # Em pasta de rede, evitamos Path.replace().
    shutil.copy2(
        caminho_temporario,
        caminho
    )

    caminho_temporario.unlink(
        missing_ok=True
    )

    return {
        "peg": peg,
        "quantidade_caixa": quantidade_caixa,
        "linha": linha_encontrada,
        "nova_linha": nova_linha,
        "mensagem": (
            "PEG adicionado à planilha."
            if nova_linha
            else "PEG atualizado na planilha."
        ),
    }

def normalizar_codigo_planilha(valor):
    if pd.isna(valor):
        return ""

    texto = str(valor).strip().upper()

    if texto.endswith(".0") and texto[:-2].isdigit():
        texto = texto[:-2]

    return texto